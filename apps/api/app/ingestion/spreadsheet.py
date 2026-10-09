from __future__ import annotations

import csv
from collections.abc import Iterator
from pathlib import Path
from zipfile import ZipFile


def parse_spreadsheet(
    path: str, suffix: str, limits: tuple[int, int, int, int, int]
) -> tuple[list[dict[str, str]], list[str]]:
    maximum_rows, maximum_columns, maximum_cell, expanded_limit, text_limit = limits
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        with ZipFile(path) as archive:
            files = archive.infolist()
            if (
                len(files) > 1000
                or sum(item.file_size for item in files) > expanded_limit
                or any(item.flag_bits & 1 for item in files)
            ):
                raise ValueError("Spreadsheet expands beyond the supported size")
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if len(workbook.sheetnames) != 1:
                raise ValueError("Multiple worksheets require splitting")
            sheet = workbook.active
            if sheet is None or (sheet.max_column or 0) > maximum_columns:
                raise ValueError("Too many spreadsheet columns")
            dimensionless = sheet.max_column is None
            bounded = sheet.iter_rows(
                max_row=maximum_rows + 1,
                max_col=sheet.max_column or maximum_columns + 1,
                values_only=True,
            )

            def values() -> Iterator[tuple[object, ...]]:
                for row in bounded:
                    if dimensionless:
                        while row and row[-1] is None:
                            row = row[:-1]
                    yield row

            return collect(
                values(),
                maximum_rows,
                maximum_columns,
                maximum_cell,
                text_limit,
            )
        finally:
            workbook.close()
    with Path(path).open(encoding="utf-8-sig", newline="") as source:
        csv.field_size_limit(maximum_cell)
        return collect(csv.reader(source), maximum_rows, maximum_columns, maximum_cell, text_limit)


def collect(
    iterator: object, maximum_rows: int, maximum_columns: int, maximum_cell: int, text_limit: int
) -> tuple[list[dict[str, str]], list[str]]:
    from collections.abc import Iterable
    from typing import cast

    rows = iter(cast(Iterable[tuple[object, ...]], iterator))
    header = next(rows, None)
    if header is None:
        raise ValueError("Spreadsheet is empty")
    headers = [
        str(value).strip() if value is not None and str(value).strip() else f"Unnamed: {index}"
        for index, value in enumerate(header)
    ]
    if len(headers) > maximum_columns or any(len(value) > maximum_cell for value in headers):
        raise ValueError("Invalid or oversized spreadsheet headers")
    reserved = set(headers)
    seen: set[str] = set()
    for index, column_header in enumerate(headers):
        name = column_header
        suffix = 1
        if name in seen:
            while True:
                name = f"{column_header}.{suffix}"
                suffix += 1
                if name not in reserved and name not in seen:
                    break
        headers[index] = name
        seen.add(name)
    result = []
    total = sum(map(len, headers))
    for row in rows:
        if not row:
            continue
        if len(row) > maximum_columns or len(row) > len(headers):
            raise ValueError("Too many spreadsheet columns")
        values = ["" if value is None else str(value) for value in row]
        if any(len(value) > maximum_cell for value in values):
            raise ValueError("Oversized cell")
        total += sum(map(len, values))
        if total > text_limit:
            raise ValueError("Spreadsheet text exceeds the supported size")
        result.append(dict(zip(headers, values + [""] * (len(headers) - len(values)), strict=True)))
        if len(result) >= maximum_rows:
            break
    return result, headers


def parse_worker(
    connection: object, path: str, suffix: str, limits: tuple[int, int, int, int, int]
) -> None:
    import resource
    from multiprocessing.connection import Connection
    from typing import cast

    pipe = cast(Connection, connection)
    try:
        resource.setrlimit(resource.RLIMIT_AS, (268435456, 268435456))
        resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
        pipe.send((True, parse_spreadsheet(path, suffix, limits)))
    except Exception:
        pipe.send((False, "Unreadable or oversized spreadsheet"))
    finally:
        pipe.close()
